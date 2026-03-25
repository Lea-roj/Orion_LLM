import os
import re
import html
from docx import Document
import torch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from transformers import M2M100ForConditionalGeneration, M2M100Tokenizer
import time

INPUT_FOLDER = "documents"
OUTPUT_FOLDER = "translated_pdf"

os.makedirs(OUTPUT_FOLDER, exist_ok=True)

print("Loading translation model...")
device = "cuda" if torch.cuda.is_available() else "cpu"

MODEL_NAME = "facebook/m2m100_418M"

tokenizer = M2M100Tokenizer.from_pretrained(MODEL_NAME)
model = M2M100ForConditionalGeneration.from_pretrained(MODEL_NAME).to(device)
if device == "cuda":
    model = model.half()

model.eval()

tokenizer.src_lang = "sl"
target_lang = tokenizer.get_lang_id("en")

def read_docx(path):
    doc = Document(path)
    paragraphs = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)
    return paragraphs


def split_text(text, max_len=400):
    text = text.replace("roj.", "roj")
    text = text.replace("drž.", "drž")
    text = text.replace("št.", "št")

    sentences = re.split(r'(?<=[.!?])\s+', text)

    chunks = []
    current = ""

    for s in sentences:
        if len(current) + len(s) < max_len:
            current += " " + s
        else:
            chunks.append(current.strip())
            current = s

    if current:
        chunks.append(current.strip())

    return chunks


def translate(text):

    inputs = tokenizer(
        text,
        return_tensors="pt",
        truncation=True,
        max_length=1024
    ).to(device)

    with torch.no_grad():
        tokens = model.generate(
            **inputs,
            forced_bos_token_id=target_lang,
            num_beams=2,
            repetition_penalty=1.1,
            max_new_tokens=512,
        )

    return tokenizer.decode(tokens[0], skip_special_tokens=True)


def translate_batch(texts):

    inputs = tokenizer(
        texts,
        return_tensors="pt",
        padding=True,
        truncation=True,
        max_length=1024
    ).to(device)

    with torch.no_grad():
        tokens = model.generate(
            **inputs,
            forced_bos_token_id=target_lang,
            num_beams=2,
            repetition_penalty=1.1,
            max_new_tokens=512
        )

    return tokenizer.batch_decode(tokens, skip_special_tokens=True)


def translate_long_text(text):

    chunks = split_text(text, max_len=600)

    translated_chunks = []

    batch_size = 16

    for i in range(0, len(chunks), batch_size):
        batch = chunks[i:i+batch_size]
        translated = translate_batch(batch)
        translated_chunks.extend(translated)

    return "\n".join(translated_chunks)


def save_pdf(paragraphs, output_path):
    styles = getSampleStyleSheet()
    normal = styles["Normal"]
    story = []

    for p in paragraphs:
        safe_text = html.escape(p).replace("\n", "<br/>")
        story.append(Paragraph(safe_text, normal))
        story.append(Spacer(1, 10))

    pdf = SimpleDocTemplate(output_path)
    pdf.build(story)


def main():
    print("Starting translation pipeline...")

    start_total = time.time()

    files = [f for f in os.listdir(INPUT_FOLDER) if f.endswith(".docx")]

    if not files:
        print("No .docx files found in input folder.")
        return

    for file_name in files:
        file_start = time.time()

        input_path = os.path.join(INPUT_FOLDER, file_name)
        base_name = os.path.splitext(file_name)[0]

        pdf_output_path = os.path.join(OUTPUT_FOLDER, base_name + ".pdf")
        docx_output_path = os.path.join(OUTPUT_FOLDER, base_name + "_translated.docx")

        print(f"\nProcessing: {file_name}")

        try:
            paragraphs = read_docx(input_path)

            translated_paragraphs = []
            for i, p in enumerate(paragraphs):
                print(f"Translating paragraph {i+1}/{len(paragraphs)}")
                translated = translate_long_text(p)
                translated_paragraphs.append(translated)

            save_pdf(translated_paragraphs, pdf_output_path)
            save_docx(translated_paragraphs, docx_output_path)

            file_time = time.time() - file_start
            print(f"Saved translated PDF: {pdf_output_path}")
            print(f"Saved translated DOCX: {docx_output_path}")
            print(f"Time for {file_name}: {file_time:.2f} seconds")

        except Exception as e:
            print(f"Error processing {file_name}: {e}")

    total_time = time.time() - start_total
    print(f"\nAll files processed in {total_time:.2f} seconds.")


def save_docx(paragraphs, output_path):
    doc = Document()

    for p in paragraphs:
        doc.add_paragraph(p)

    doc.save(output_path)

if __name__ == "__main__":
    main()